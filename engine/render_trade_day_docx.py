"""
Renders the printable Trade Day pairing sheet (Word doc) from
generate_trade_day_rotation.py -- same visual style as the existing
Sponsorship Negotiation rotation doc (negotiation-roleplay/
Rotation_Schedule_20teams.docx) since Trade Day is a sibling in-class
pairing exercise.

Usage:
    python3 engine/render_trade_day_docx.py --n-teams 18
"""
import argparse
import json
import os
import sys

from docx import Document

sys.path.insert(0, os.path.dirname(__file__))
from generate_trade_day_rotation import build_trade_day

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def student_labels(n_teams):
    """'Owner Name (Team N)' labels so students can find themselves by their
    real name, not just the placeholder team name -- falls back to the
    placeholder alone if league_config.json doesn't have n_teams teams or an
    owner isn't filled in yet."""
    config_path = os.path.join(BASE, "data", "league_config.json")
    with open(config_path) as f:
        config = json.load(f)
    if len(config["teams"]) != n_teams:
        return None
    return [
        f"{t['owner']} ({t['name']})" if t.get("owner") else t["name"]
        for t in config["teams"]
    ]


def render(n_teams):
    team_names = student_labels(n_teams)
    trade_day = build_trade_day(n_teams, team_names)

    with open(os.path.join(BASE, "data", "deals.json")) as f:
        deals = json.load(f)
    tr = deals["trade_rules"]
    # split/deadline round numbers live only in this narrative string in
    # deals.json -- reusing it here means this doc can't drift out of sync
    # with the rulebook the way the old hardcoded "Round 10"/"Round 15" text did.

    doc = Document()
    doc.add_heading(f"Mandatory Mid-Season Trade Day — Rotation Schedule ({n_teams} teams)", level=0)
    doc.add_paragraph(tr["mid_season_trade_day"])

    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Meeting"
    hdr[1].text = "Paired Teams"
    for i, p in enumerate(trade_day["pairs"], start=1):
        row = table.add_row().cells
        row[0].text = f"Meeting {i}"
        row[1].text = f"{p['team_a']} & {p['team_b']}"

    if trade_day["bye"]:
        doc.add_heading("Bye", level=1)
        doc.add_paragraph(
            f"{', '.join(trade_day['bye'])} draws the bye this round (odd team count) "
            f"-- the instructor fills in as trade partner so this team still completes "
            f"a mandatory trade, same convention as the Sponsorship Negotiation rotation's bye."
        )

    out_dir = os.path.join(BASE, "negotiation-roleplay")
    out_path = os.path.join(out_dir, f"Trade_Day_Rotation_{n_teams}teams.docx")
    doc.save(out_path)
    print(f"Saved -> {out_path}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--n-teams", type=int, required=True)
    args = ap.parse_args()
    render(args.n_teams)
